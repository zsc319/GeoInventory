from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app import create_app
from geo_inventory.format_conversion import (
    PERFORATION_STANDARD_HEADERS,
    preview_petrel_perforations,
    write_petrel_perforation_workbook,
)


def _petrel_perforation_file(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Petrel export"
    sheet.append(["UNITS", "FIELD"])
    sheet.append([])
    sheet.append(["WELLNAME", "ET-101"])
    sheet.append(["26/4/1982", "perforation", 644.685, 664.370, 1.640, 0, 0, 0])
    sheet.append(["27/4/1982", "squeeze", 700.0, 702.0, None, None, None, None])
    sheet.append([])
    sheet.append(["WELLNAME", "ET-102"])
    sheet.append([date(1983, 7, 21), "perforation", 661.0, 669.0, 8.0, 0, 0, 0])
    workbook.save(path)
    workbook.close()


def test_preview_petrel_perforations_extracts_grouped_wells_and_drops_zero_tail(tmp_path):
    source = tmp_path / "petrel_perforations.xlsx"
    _petrel_perforation_file(source)

    preview = preview_petrel_perforations(source)

    assert preview["headers"] == PERFORATION_STANDARD_HEADERS
    assert preview["summary"]["well_count"] == 2
    assert preview["summary"]["record_count"] == 3
    assert preview["summary"]["event_types"] == [
        {"name": "perforation", "count": 2},
        {"name": "squeeze", "count": 1},
    ]
    assert [row["source_column"] for row in preview["summary"]["ignored_zero_columns"]] == [
        "第 6 列", "第 7 列", "第 8 列"
    ]
    first = preview["rows"][0]
    assert first["WELL NAME"] == "ET-101"
    assert first["TIME"] == "1982-04-26"
    assert first["COMPLETION TYPE"] == "perforation"
    assert first["TOP"] == 644.685
    assert first["BASE"] == 664.37
    assert round(first["THICKNESS"], 3) == 19.685
    assert first["WELLBORE DIAMETER"] == 1.64
    assert preview["rows"][1]["THICKNESS"] == 2
    assert preview["rows"][1]["WELLBORE DIAMETER"] is None


def test_perforation_standard_workbook_has_only_requested_columns(tmp_path):
    source = tmp_path / "petrel_perforations.xlsx"
    destination = tmp_path / "standard_perforations.xlsx"
    _petrel_perforation_file(source)

    result = write_petrel_perforation_workbook(source, destination)
    workbook = load_workbook(destination, data_only=True)
    sheet = workbook["Perforations"]

    assert result == {"records": 3, "wells": 2}
    assert [cell.value for cell in sheet[1]] == PERFORATION_STANDARD_HEADERS
    assert sheet.max_column == 7
    assert sheet.max_row == 4
    assert sheet["A2"].value == "ET-101"
    assert sheet["C3"].value == "squeeze"
    assert sheet["F4"].value == 8
    assert sheet["G2"].value == 1.64
    assert sheet.freeze_panes == "A2"
    workbook.close()


def test_perforation_batch_export_writes_unique_files_and_manifest(tmp_path):
    first = tmp_path / "ET perforations.xlsx"
    second = tmp_path / "ET perforations copy.xlsx"
    output_dir = tmp_path / "converted"
    output_dir.mkdir()
    _petrel_perforation_file(first)
    _petrel_perforation_file(second)
    app = create_app({"TESTING": True, "DATABASE": str(tmp_path / "inventory.sqlite")})

    response = app.test_client().post(
        "/api/format-conversion/perforation/export-batch",
        json={"paths": [str(first), str(second)], "output_dir": str(output_dir)},
    )

    assert response.status_code == 200
    payload = response.get_json()
    assert len(payload["converted"]) == 2
    assert not payload["failed"]
    assert (output_dir / payload["converted"][0]["output_file"]).is_file()
    assert (output_dir / payload["converted"][1]["output_file"]).is_file()
    assert (output_dir / Path(payload["manifest"]).name).is_file()
