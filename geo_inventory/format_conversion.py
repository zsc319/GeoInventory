from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any


PERFORATION_STANDARD_HEADERS = [
    "WELL NAME",
    "TIME",
    "COMPLETION TYPE",
    "TOP",
    "BASE",
    "THICKNESS",
    "WELLBORE DIAMETER",
]


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _marker(value: Any) -> str:
    return "".join(char for char in _text(value).upper() if char.isalnum())


def _number(value: Any) -> float | int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def _normalise_time(value: Any) -> tuple[str, date | datetime | None]:
    if isinstance(value, datetime):
        return value.date().isoformat(), value
    if isinstance(value, date):
        return value.isoformat(), value
    raw = _text(value)
    if not raw:
        return "", None
    for pattern in ("%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(raw, pattern).date()
            return parsed.isoformat(), parsed
        except ValueError:
            continue
    return raw, None


def _zero_or_blank(value: Any) -> bool:
    if value is None or _text(value) == "":
        return True
    number = _number(value)
    return number == 0 if number is not None else False


def _load_workbook(path: Path):
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("射孔数据整理当前支持 .xlsx 或 .xlsm 文件")
    if not path.is_file():
        raise ValueError("所选 Excel 文件不存在或无法访问")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # packaged build carries openpyxl; retain a useful dev error
        raise RuntimeError("当前运行环境缺少 Excel 读取组件 openpyxl") from exc
    return load_workbook(path, read_only=True, data_only=True)


def read_petrel_perforations(path: str | Path) -> dict[str, Any]:
    """Read Petrel's grouped perforation export without changing the source file.

    Expected blocks start with ``WELLNAME`` in column A and the well name in
    column B. Detail rows use A:E as time, completion type, top, base and
    wellbore diameter. Thickness is always calculated as ``BASE - TOP``. In
    this Petrel event layout, column E is not an interval thickness. Columns
    F:H are detected only for reporting; they are never carried into the
    standard delivery table.
    """

    source = Path(path).expanduser().resolve()
    workbook = _load_workbook(source)
    records: list[dict[str, Any]] = []
    event_counts: Counter[str] = Counter()
    well_names: set[str] = set()
    detail_rows = 0
    skipped_rows = 0
    non_positive_thickness_count = 0
    trailing_values: list[list[Any]] = [[], [], []]
    source_layout: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            current_well = ""
            sheet_headers = 0
            sheet_records = 0
            for row_number, raw_values in enumerate(sheet.iter_rows(values_only=True), 1):
                values = list(raw_values)
                if _marker(values[0] if values else None) == "WELLNAME":
                    current_well = _text(values[1] if len(values) > 1 else None)
                    if current_well:
                        well_names.add(current_well)
                        sheet_headers += 1
                    continue
                if not current_well or not any(value is not None and _text(value) for value in values[:5]):
                    continue

                detail_rows += 1
                time_text, time_value = _normalise_time(values[0] if len(values) > 0 else None)
                completion_type = _text(values[1] if len(values) > 1 else None)
                top = _number(values[2] if len(values) > 2 else None)
                base = _number(values[3] if len(values) > 3 else None)
                wellbore_diameter = _number(values[4] if len(values) > 4 else None)
                thickness = (base - top) if base is not None and top is not None else None
                for index in range(3):
                    trailing_values[index].append(values[5 + index] if len(values) > 5 + index else None)

                if not completion_type or top is None or base is None:
                    skipped_rows += 1
                    continue
                if thickness is not None and thickness <= 0:
                    non_positive_thickness_count += 1
                record = {
                    "WELL NAME": current_well,
                    "TIME": time_text,
                    "COMPLETION TYPE": completion_type,
                    "TOP": top,
                    "BASE": base,
                    "THICKNESS": thickness,
                    "WELLBORE DIAMETER": wellbore_diameter,
                    "_time_value": time_value,
                    "_sheet": sheet.title,
                    "_row": row_number,
                }
                records.append(record)
                event_counts[completion_type] += 1
                sheet_records += 1
            source_layout.append({"sheet": sheet.title, "well_headers": sheet_headers, "records": sheet_records})
    finally:
        workbook.close()

    if not well_names:
        raise ValueError("未识别到 Petrel 射孔块头：需要 A 列为 WELLNAME、B 列为井名")
    if not records:
        raise ValueError("已识别 WELLNAME，但未找到包含完井事件和深度段的射孔明细")

    ignored_zero_columns = [
        {"source_column": f"第 {index + 6} 列", "reason": "明细记录中均为 0 或空值，未纳入标准表"}
        for index, values in enumerate(trailing_values)
        if values and all(_zero_or_blank(value) for value in values)
    ]
    return {
        "source_file": source.name,
        "source_path": str(source),
        "headers": PERFORATION_STANDARD_HEADERS,
        "records": records,
        "summary": {
            "well_count": len(well_names),
            "record_count": len(records),
            "detail_rows": detail_rows,
            "skipped_rows": skipped_rows,
            "non_positive_thickness_count": non_positive_thickness_count,
            "event_types": [
                {"name": name, "count": count}
                for name, count in sorted(event_counts.items(), key=lambda item: (-item[1], item[0].upper()))
            ],
            "source_layout": source_layout,
            "ignored_zero_columns": ignored_zero_columns,
            "input_mapping": [
                {"source": "块头行：A 列 WELLNAME，B 列井名", "target": "WELL NAME"},
                {"source": "明细行 A 列", "target": "TIME"},
                {"source": "明细行 B 列", "target": "COMPLETION TYPE"},
                {"source": "明细行 C / D 列", "target": "TOP / BASE"},
                {"source": "明细行 C / D 列相减（BASE − TOP）", "target": "THICKNESS"},
                {"source": "明细行 E 列（Petrel Wellbore Diameter）", "target": "WELLBORE DIAMETER"},
            ],
        },
    }


def preview_petrel_perforations(path: str | Path, limit: int = 100) -> dict[str, Any]:
    parsed = read_petrel_perforations(path)
    rows = []
    for record in parsed["records"][:max(1, min(int(limit), 500))]:
        rows.append({header: record.get(header) for header in PERFORATION_STANDARD_HEADERS})
    return {
        "source_file": parsed["source_file"],
        "headers": parsed["headers"],
        "summary": parsed["summary"],
        "rows": rows,
        "truncated": len(parsed["records"]) > len(rows),
    }


def write_petrel_perforation_workbook(path: str | Path, destination: str | Path) -> dict[str, Any]:
    parsed = read_petrel_perforations(path)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("当前运行环境缺少 Excel 写入组件 openpyxl") from exc

    output = Path(destination)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Perforations"
    sheet.append(PERFORATION_STANDARD_HEADERS)
    for record in parsed["records"]:
        sheet.append([
            record["WELL NAME"],
            record["_time_value"] or record["TIME"],
            record["COMPLETION TYPE"],
            record["TOP"],
            record["BASE"],
            record["THICKNESS"],
            record["WELLBORE DIAMETER"],
        ])
    header = sheet[1]
    for cell in header:
        cell.fill = PatternFill("solid", fgColor="16735F")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in sheet["B"][1:]:
        if isinstance(cell.value, (date, datetime)):
            cell.number_format = "yyyy-mm-dd"
    for column in ("D", "E", "F", "G"):
        for cell in sheet[column][1:]:
            cell.number_format = "0.000"
    for column, width in {"A": 28, "B": 14, "C": 22, "D": 14, "E": 14, "F": 14, "G": 20}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{len(parsed['records']) + 1}"
    workbook.save(output)
    workbook.close()
    return {"records": len(parsed["records"]), "wells": parsed["summary"]["well_count"]}
