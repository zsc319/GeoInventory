from __future__ import annotations

import csv
import json
import math
import re
import sqlite3
import struct
import unicodedata
import zlib
from collections import defaultdict
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, BinaryIO, Callable, Iterator

from .db import Database, utcnow
from .geometry import convex_hull, polygon_area, ring_area
from .ofm_mdb import import_mdb
from .trajectory import minimum_curvature


SOURCE_PRIORITY = {
    "well_head": 90,
    "deviation": 76,
    "las": 62,
    "core": 54,
    "interpretation": 48,
    "production": 44,
}

FIELD_ALIASES = {
    "well": {"WELL", "WELLNAME", "WELLID", "BOREHOLE", "HOLENAME", "井名", "井号"},
    "uwi": {"UWI", "API", "WELLUWI", "UNIQUEWELLID", "统一井号"},
    "x": ("SURFACEX", "WELLHEADX", "地面X", "井口X", "X", "EASTING", "XCOORD", "XCOORDINATE", "LONGITUDE", "LON", "经度", "横坐标"),
    "y": ("SURFACEY", "WELLHEADY", "地面Y", "井口Y", "Y", "NORTHING", "YCOORD", "YCOORDINATE", "LATITUDE", "LAT", "纬度", "纵坐标"),
    "base_x": {"BASEX", "BOTTOMX", "TDX", "井底X"},
    "base_y": {"BASEY", "BOTTOMY", "TDY", "井底Y"},
    "kb": {"KB", "KBELEVATION", "ELEVATION", "DERRICKELEVATION", "DATUMELEVATION", "补心海拔", "井口海拔"},
    "td": {"TD", "TOTALDEPTH", "WELLDEPTH", "完钻井深", "总井深"},
    "md": {"MD", "MEASUREDDEPTH", "DEPTH", "DEPT", "井深", "斜深"},
    "inc": {"INC", "INCL", "INCLINATION", "ANGLE", "井斜", "井斜角"},
    "azi": {"AZI", "AZIMUTH", "AZM", "方位", "方位角"},
    "top": {"TOP", "TOPMD", "FROM", "STARTDEPTH", "顶深", "顶界"},
    "base": {"BASE", "BOTTOM", "BOT", "BASEMD", "TO", "ENDDEPTH", "底深", "底界"},
    "name": {"NAME", "POLYGON", "POLYGONNAME", "区块", "多边形", "名称"},
    "alias": {"ALIAS", "ALIASNAME", "SHORTNAME", "别名", "缩写井名"},
    "canonical": {"CANONICAL", "CANONICALNAME", "FULLNAME", "标准井名", "正式井名"},
    "version": {"VERSION", "INTERPRETATIONVERSION", "REVISION", "REV", "版本"},
    "batch": {"BATCH", "BATCHNAME", "批次"},
    "interval_name": {"INTERVAL", "INTERVALNAME", "ZONE", "PERFORATION", "射孔段", "层段"},
    "status": {"STATUS", "STATE", "OPENSTATUS", "状态", "开关状态"},
    "event_date": {"DATE", "EVENTDATE", "STARTDATE", "PERFORATIONDATE", "日期", "射孔日期"},
    "production_month": {"PRODUCTIONMONTH", "PRODMONTH", "YEARMONTH", "MONTH", "PRODDATE", "DATE", "生产月份", "生产日期", "年月"},
    "onstream_date": {"ONSTREAMDATE", "FIRSTPRODUCTIONDATE", "STARTPRODUCTION", "COMMENCEMENTDATE", "投产时间", "投产日期"},
    "days_on": {"DAYSON", "PRODDAYS", "DAYSPRODUCING", "生产天数", "开井天数"},
    "liquid_rate": {"LIQUIDRATE", "LIQRATE", "QL", "Q液", "日产液", "日产液量"},
    "oil_rate": {"OILRATE", "QO", "OILPRODRATE", "日产油", "日产油量"},
    "water_rate": {"WATERRATE", "QW", "WATERPRODRATE", "日产水", "日产水量"},
    "water_cut": {"WATERCUT", "WCUT", "WCT", "FW", "含水率", "含水"},
    "monthly_oil": {"MONTHLYOIL", "OILMONTH", "OILVOLUME", "MOIL", "月产油", "月产油量"},
    "monthly_water": {"MONTHLYWATER", "WATERMONTH", "WATERVOLUME", "MWATER", "月产水", "月产水量"},
    "cumulative_oil": {"CUMULATIVEOIL", "CUMOIL", "OILCUM", "NP", "累产油", "累计产油"},
    "cumulative_water": {"CUMULATIVEWATER", "CUMWATER", "WATERCUM", "WP", "累产水", "累计产水"},
    "pressure": {"PRESSURE", "BHP", "THP", "WHP", "RESERVOIRPRESSURE", "压力", "井底流压", "油压", "套压"},
    "pressure_type": {"PRESSURETYPE", "PRESSUREKIND", "压力类型"},
    "event_type": {"EVENT", "EVENTTYPE", "OPERATION", "开关井事件", "事件类型"},
}


def normalize_header(value: Any) -> str:
    value = unicodedata.normalize("NFKC", str(value or "")).strip().upper()
    return re.sub(r"[^0-9A-Z\u4e00-\u9fff]+", "", value)


def normalize_well_name(value: Any) -> str:
    return normalize_header(value)


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return None if text == "" or text.upper() in {"NULL", "NONE", "N/A", "NA", "NAN", "-999.25"} else text


def number(value: Any) -> float | None:
    text = clean_text(value)
    if text is None:
        return None
    try:
        result = float(text.replace(",", ""))
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def field_map(headers: list[str]) -> dict[str, str]:
    normalized = {normalize_header(h): h for h in headers}
    mapped: dict[str, str] = {}
    for logical, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                mapped[logical] = normalized[alias]
                break
    return mapped


def detect_type(path: str | Path, headers: list[str] | None = None) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    name = path.stem.lower()
    if suffix == ".las":
        return "las"
    if suffix in {".mdb", ".accdb"}:
        return "production"
    if suffix in {".sgy", ".segy", ".seg-y", ".zgy"}:
        return "seismic"
    if suffix in {".geojson", ".json"}:
        return "polygon"
    if headers:
        mapped = field_map(headers)
        normalized = {normalize_header(h) for h in headers}
        production_name = any(token in name for token in ("production", "perfor", "completion", "ofm", "prod", "生产", "射孔"))
        production_fields = {"production_month", "onstream_date", "days_on", "liquid_rate", "oil_rate", "water_rate", "water_cut", "monthly_oil", "monthly_water", "cumulative_oil", "cumulative_water", "pressure", "event_type"}
        if production_name and "well" in mapped and ({"top", "base"} <= mapped.keys() or bool(production_fields & mapped.keys())):
            return "production"
        if {"alias", "canonical"} <= mapped.keys():
            return "alias"
        if {"well", "md", "inc", "azi"} <= mapped.keys():
            return "deviation"
        if "well" in mapped and ({"top", "base"} & mapped.keys() or any(x in normalized for x in {"PHI", "POR", "PERM", "SW", "FACIES", "LITHOLOGY"})):
            return "interpretation"
        if {"x", "y", "name"} <= mapped.keys() and "well" not in mapped:
            return "polygon"
        if "well" in mapped:
            return "well_head"
    if "dev" in name or "deviation" in name or "轨迹" in name or "井斜" in name:
        return "deviation"
    if "core" in name or "岩心" in name:
        return "core"
    if "top" in name or "interpret" in name or "解释" in name:
        return "interpretation"
    return "well_head"


def read_tabular(path: str | Path) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("读取 Excel 需要 openpyxl：pip install openpyxl") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if not raw_headers:
            workbook.close()
            return [], iter(())
        headers = [str(v).strip() if v is not None else f"COLUMN_{i + 1}" for i, v in enumerate(raw_headers)]

        def iterator() -> Iterator[tuple[int, dict[str, Any]]]:
            try:
                for row_no, values in enumerate(rows, 2):
                    yield row_no, {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
            finally:
                workbook.close()

        return headers, iterator()

    with path.open("rb") as probe:
        raw_sample = probe.read(65536)
    encoding_name = None
    for encoding in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            raw_sample.decode(encoding)
            encoding_name = encoding
            break
        except UnicodeDecodeError:
            continue
    if encoding_name is None:
        raise ValueError("无法识别文本编码")
    text_handle = path.open("r", encoding=encoding_name, errors="replace", newline="")
    sample = text_handle.read(8192)
    text_handle.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(text_handle, dialect=dialect)
    headers = [str(h).strip() for h in (reader.fieldnames or [])]

    def iterator() -> Iterator[tuple[int, dict[str, Any]]]:
        try:
            for row_no, row in enumerate(reader, 2):
                yield row_no, {str(k).strip(): v for k, v in row.items() if k is not None}
        finally:
            text_handle.close()

    return headers, iterator()


@dataclass
class ImportResult:
    source_id: int
    filename: str
    data_type: str
    records: int = 0
    warnings: list[str] = field(default_factory=list)
    details: dict[str, Any] = field(default_factory=dict)


def _report_progress(callback: Callable[[float, str], None] | None, fraction: float, stage: str) -> None:
    if callback:
        callback(max(0.0, min(1.0, fraction)), stage)


class ImportService:
    def __init__(self, database: Database):
        self.database = database

    def import_path(
        self,
        path: str | Path,
        data_type: str = "auto",
        batch: str | None = None,
        version: str | None = None,
        crs: str | None = None,
        progress_callback: Callable[[float, str], None] | None = None,
    ) -> ImportResult:
        path = Path(path).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"文件不存在：{path}")
        if data_type == "auto":
            _report_progress(progress_callback, 0.02, "识别文件类型")
            headers = None
            if path.suffix.lower() in {".csv", ".txt", ".tsv", ".xlsx", ".xlsm"}:
                headers, _ = read_tabular(path)
            data_type = detect_type(path, headers)
        supported = {"well_head", "deviation", "las", "interpretation", "core", "production", "seismic", "polygon", "alias"}
        if data_type not in supported:
            raise ValueError(f"不支持的数据类型：{data_type}")

        with self.database.connect() as conn:
            _report_progress(progress_callback, 0.08, f"准备导入 {data_type}")
            cursor = conn.execute(
                "INSERT INTO sources(filename,file_path,data_type,batch,version,crs,imported_at) VALUES(?,?,?,?,?,?,?)",
                (path.name, str(path), data_type, clean_text(batch), clean_text(version), clean_text(crs), utcnow()),
            )
            source_id = cursor.lastrowid
            result = ImportResult(source_id, path.name, data_type)
            try:
                handler = getattr(self, f"_import_{data_type}")
                handler(conn, path, result, crs=clean_text(crs), version=clean_text(version), progress_callback=progress_callback)
                conn.execute(
                    "UPDATE sources SET status='ready',record_count=?,warning=?,metadata_json=? WHERE id=?",
                    (result.records, "\n".join(result.warnings) or None, Database.json(result.details), source_id),
                )
                conn.commit()
                _report_progress(progress_callback, 1.0, "写入数据库完成")
                return result
            except Exception as exc:
                conn.rollback()
                # Preserve a failure audit record in a separate transaction.
                conn.execute(
                    "INSERT INTO sources(id,filename,file_path,data_type,batch,version,crs,imported_at,status,warning) VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (source_id, path.name, str(path), data_type, clean_text(batch), clean_text(version), clean_text(crs), utcnow(), "failed", str(exc)),
                )
                conn.commit()
                raise

    def _resolve_well(
        self, conn: sqlite3.Connection, raw_name: Any, uwi: Any = None
    ) -> int:
        raw = clean_text(raw_name)
        if not raw:
            raise ValueError("发现空井名")
        key = normalize_well_name(raw)
        rule = conn.execute("SELECT canonical_key,canonical_name FROM alias_rules WHERE alias_key=?", (key,)).fetchone()
        canonical_key = rule["canonical_key"] if rule else key
        canonical_name = rule["canonical_name"] if rule else raw
        uwi_text = clean_text(uwi)
        row = conn.execute("SELECT id FROM wells WHERE normalized_key=?", (canonical_key,)).fetchone()
        if not row and uwi_text:
            row = conn.execute("SELECT id FROM wells WHERE uwi=?", (uwi_text,)).fetchone()
        if row:
            return int(row["id"])
        cursor = conn.execute(
            "INSERT INTO wells(canonical_name,normalized_key,uwi,created_at) VALUES(?,?,?,?)",
            (canonical_name, canonical_key, uwi_text, utcnow()),
        )
        return int(cursor.lastrowid)

    def _add_occurrence(
        self,
        conn: sqlite3.Connection,
        result: ImportResult,
        raw_name: Any,
        row_no: int | None,
        values: dict[str, Any],
    ) -> int:
        well_id = self._resolve_well(conn, raw_name, values.get("uwi"))
        conn.execute(
            """INSERT INTO well_sources(
                well_id,source_id,raw_name,normalized_name,row_no,uwi,x,y,crs,kb_elevation,total_depth,attributes_json
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                well_id, result.source_id, str(raw_name).strip(), normalize_well_name(raw_name), row_no,
                clean_text(values.get("uwi")), number(values.get("x")), number(values.get("y")), clean_text(values.get("crs")),
                number(values.get("kb")), number(values.get("td")), Database.json(values.get("attributes", {})),
            ),
        )
        self._refresh_well(conn, well_id)
        return well_id

    def _refresh_well(self, conn: sqlite3.Connection, well_id: int) -> None:
        rows = conn.execute(
            """SELECT ws.*,s.data_type FROM well_sources ws JOIN sources s ON s.id=ws.source_id
               WHERE ws.well_id=? ORDER BY ws.id""",
            (well_id,),
        ).fetchall()
        if not rows:
            return
        ranked = sorted(rows, key=lambda row: SOURCE_PRIORITY.get(row["data_type"], 30), reverse=True)

        def best(field: str) -> Any:
            return next((row[field] for row in ranked if row[field] is not None and row[field] != ""), None)

        preferred = ranked[0]
        source_types = {row["data_type"] for row in rows}
        confidence = SOURCE_PRIORITY.get(preferred["data_type"], 30)
        if best("uwi"):
            confidence += 5
        if best("x") is not None and best("y") is not None:
            confidence += 3
        confidence += min(4, max(0, len(source_types) - 1) * 2)
        conn.execute(
            """UPDATE wells SET uwi=COALESCE(?,uwi),x=?,y=?,crs=?,kb_elevation=?,total_depth=?,
               preferred_source_id=?,confidence_score=? WHERE id=?""",
            (best("uwi"), best("x"), best("y"), best("crs"), best("kb_elevation"), best("total_depth"), preferred["source_id"], min(100, confidence), well_id),
        )

    def _import_well_head(self, conn: sqlite3.Connection, path: Path, result: ImportResult, crs: str | None, progress_callback=None, **_: Any) -> None:
        headers, rows = read_tabular(path)
        mapped = field_map(headers)
        if "well" not in mapped:
            raise ValueError("井头表缺少井名列（如 WELL/WELL_NAME/井名）")
        missing_coordinates = 0
        for row_no, row in rows:
            raw_name = row.get(mapped["well"])
            if not clean_text(raw_name):
                continue
            values = {key: row.get(header) for key, header in mapped.items() if key != "well"}
            values["crs"] = crs
            values["attributes"] = row
            if number(values.get("x")) is None or number(values.get("y")) is None:
                missing_coordinates += 1
            self._add_occurrence(conn, result, raw_name, row_no, values)
            result.records += 1
            if result.records % 256 == 0:
                _report_progress(progress_callback, 0.12 + 0.8 * result.records / max(1, len(rows)), "写入井头记录")
        result.details["field_mapping"] = mapped
        result.details["coordinate_role"] = "surface X / surface Y 用作井口二维投影；base X / base Y 作为原始井底属性保留"
        result.details["located_records"] = result.records - missing_coordinates
        result.details["missing_coordinate_records"] = missing_coordinates
        if missing_coordinates:
            result.warnings.append(f"{missing_coordinates} 口井缺少完整地面 X/Y，已入库但不会出现在二维投影中")

    def _import_deviation(self, conn: sqlite3.Connection, path: Path, result: ImportResult, crs: str | None, progress_callback=None, **_: Any) -> None:
        headers, rows = read_tabular(path)
        mapped = field_map(headers)
        missing = {"well", "md", "inc", "azi"} - mapped.keys()
        if missing:
            raise ValueError(f"井斜表缺少字段：{', '.join(sorted(missing))}")
        grouped: dict[str, list[tuple[int, dict[str, Any], dict[str, Any]]]] = defaultdict(list)
        for row_no, row in rows:
            raw_name = clean_text(row.get(mapped["well"]))
            md, inc, azi = number(row.get(mapped["md"])), number(row.get(mapped["inc"])), number(row.get(mapped["azi"]))
            if not raw_name or md is None or inc is None or azi is None:
                continue
            grouped[raw_name].append((row_no, {"md": md, "inclination": inc, "azimuth": azi}, row))
        for group_index, (raw_name, station_rows) in enumerate(grouped.items(), 1):
            first_row = station_rows[0][2]
            values = {key: first_row.get(header) for key, header in mapped.items() if key not in {"well", "md", "inc", "azi"}}
            values.update({"crs": crs, "attributes": first_row})
            well_id = self._add_occurrence(conn, result, raw_name, station_rows[0][0], values)
            kb = number(values.get("kb"))
            calculated = minimum_curvature([station for _, station, _ in station_rows], kb)
            conn.executemany(
                """INSERT INTO deviation_stations(
                   source_id,well_id,md,inclination,azimuth,tvd,northing,easting,tvdss,z_msl
                   ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                [(
                    result.source_id, well_id, row["md"], row["inclination"], row["azimuth"], row["tvd"], row["northing"], row["easting"], row.get("tvdss"), row.get("z_msl")
                ) for row in calculated],
            )
            result.records += len(calculated)
            _report_progress(progress_callback, 0.15 + 0.75 * group_index / max(1, len(grouped)), "换算 MD / TVDSS 并写入井轨迹")
        result.details.update({"wells": len(grouped), "field_mapping": mapped, "tvdss_convention": "positive_down: TVD - KB_elevation"})

    def _import_las(self, conn: sqlite3.Connection, path: Path, result: ImportResult, crs: str | None, progress_callback=None, **_: Any) -> None:
        parsed = parse_las(path, lambda value, stage: _report_progress(progress_callback, 0.1 + value * 0.72, stage))
        raw_name = parsed["well"].get("WELL", {}).get("value") or path.stem
        values = {
            "uwi": parsed["well"].get("UWI", {}).get("value") or parsed["well"].get("API", {}).get("value"),
            "kb": parsed["well"].get("ELEV", {}).get("value") or parsed["well"].get("KB", {}).get("value"),
            "td": parsed["well"].get("STOP", {}).get("value"),
            "crs": crs,
            "attributes": {key: value.get("value") for key, value in parsed["well"].items()},
        }
        well_id = self._add_occurrence(conn, result, raw_name, None, values)
        depth_curve = parsed["curves"][0] if parsed["curves"] else {"unit": None}
        cursor = conn.execute(
            """INSERT INTO las_files(source_id,well_id,start_md,stop_md,step,depth_unit,null_value,curve_count,sample_rows)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (
                result.source_id, well_id, parsed["start"], parsed["stop"], parsed["step"], depth_curve.get("unit"),
                parsed["null"], max(0, len(parsed["curves"]) - 1), parsed["sample_rows"],
            ),
        )
        las_file_id = int(cursor.lastrowid)
        for curve in parsed["curves"][1:]:
            conn.execute(
                """INSERT INTO las_curves(las_file_id,source_id,well_id,mnemonic,unit,description,sample_count,value_min,value_max)
                   VALUES(?,?,?,?,?,?,?,?,?)""",
                (
                    las_file_id, result.source_id, well_id, curve["mnemonic"], curve.get("unit"), curve.get("description"),
                    curve.get("sample_count", 0), curve.get("value_min"), curve.get("value_max"),
                ),
            )
        _report_progress(progress_callback, 0.9, "写入 LAS 曲线索引")
        result.records = parsed["sample_rows"]
        result.details.update({"well": raw_name, "curve_count": max(0, len(parsed["curves"]) - 1), "start_md": parsed["start"], "stop_md": parsed["stop"]})
        result.warnings.extend(parsed["warnings"])

    def _import_interpretation(self, conn: sqlite3.Connection, path: Path, result: ImportResult, crs: str | None, version: str | None, progress_callback=None, **_: Any) -> None:
        self._import_attribute_table(conn, path, result, crs, version, progress_callback)

    def _import_core(self, conn: sqlite3.Connection, path: Path, result: ImportResult, crs: str | None, version: str | None, progress_callback=None, **_: Any) -> None:
        self._import_attribute_table(conn, path, result, crs, version, progress_callback)

    def _import_production(self, conn: sqlite3.Connection, path: Path, result: ImportResult, progress_callback=None, **_: Any) -> None:
        """Import perforations and OFM-style monthly production records.

        Production names remain independent normalized keys. They are exposed
        as hard well evidence by the project-well combiner, while fuzzy names
        stay separate until an operator confirms the alias.
        """
        if path.suffix.lower() in {".mdb", ".accdb"}:
            conn.execute(
                "DELETE FROM sources WHERE id<>? AND data_type='production' AND lower(file_path)=lower(?)",
                (result.source_id, str(path)),
            )
            details = import_mdb(conn, path, result.source_id, progress_callback)
            result.records = int(details["total_rows"])
            result.details.update(details)
            if details.get("warning"):
                result.warnings.append(details["warning"])
            return
        headers, rows = read_tabular(path)
        mapped = field_map(headers)
        if "well" not in mapped:
            raise ValueError("生产动态表缺少井名字段（WELL / WELL_NAME）")
        has_interval = {"top", "base"} <= mapped.keys()
        dynamic_fields = {
            "production_month", "onstream_date", "days_on", "liquid_rate", "oil_rate", "water_rate",
            "water_cut", "monthly_oil", "monthly_water", "cumulative_oil", "cumulative_water",
            "pressure", "pressure_type", "event_type",
        }
        has_dynamic = bool((dynamic_fields - {"production_month", "event_type"}) & mapped.keys()) or (
            "production_month" in mapped and not has_interval
        )
        if not has_interval and not has_dynamic:
            raise ValueError("生产表需包含射孔顶/底 MD，或生产月份、日产量、累产量、含水率、压力等动态字段")
        numeric_fields = (
            "days_on", "liquid_rate", "oil_rate", "water_rate", "water_cut", "monthly_oil",
            "monthly_water", "cumulative_oil", "cumulative_water", "pressure",
        )
        previous_status: dict[str, str] = {}
        onstream_seen: set[str] = set()
        interval_count = monthly_count = event_count = source_rows = 0
        for index, (row_no, row) in enumerate(rows, 1):
            well_name = clean_text(row.get(mapped["well"]))
            if not well_name:
                continue
            key = normalize_well_name(well_name)
            status = clean_text(row.get(mapped["status"])) if "status" in mapped else None
            date_value = clean_text(row.get(mapped["production_month"])) if "production_month" in mapped else None
            event_date = clean_text(row.get(mapped["event_date"])) if "event_date" in mapped else date_value
            inserted = False
            if has_interval:
                top = number(row.get(mapped["top"]))
                base = number(row.get(mapped["base"]))
                if top is not None and base is not None:
                    if base < top:
                        top, base = base, top
                    conn.execute(
                        """INSERT INTO production_intervals(
                           source_id,well_key,well_name,top_md,base_md,interval_name,status,event_date,metadata_json
                           ) VALUES(?,?,?,?,?,?,?,?,?)""",
                        (
                            result.source_id, key, well_name, top, base,
                            clean_text(row.get(mapped["interval_name"])) if "interval_name" in mapped else None,
                            status, event_date, Database.json(row),
                        ),
                    )
                    interval_count += 1
                    inserted = True
            values = {field: number(row.get(mapped[field])) if field in mapped else None for field in numeric_fields}
            if has_dynamic and (date_value or any(value is not None for value in values.values()) or status):
                conn.execute(
                    """INSERT INTO production_monthly(
                       source_id,well_key,well_name,production_month,days_on,liquid_rate,oil_rate,water_rate,
                       water_cut,monthly_oil,monthly_water,cumulative_oil,cumulative_water,pressure,
                       pressure_type,status,metadata_json
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        result.source_id, key, well_name, date_value, values["days_on"], values["liquid_rate"],
                        values["oil_rate"], values["water_rate"], values["water_cut"], values["monthly_oil"],
                        values["monthly_water"], values["cumulative_oil"], values["cumulative_water"],
                        values["pressure"], clean_text(row.get(mapped["pressure_type"])) if "pressure_type" in mapped else None,
                        status, Database.json(row),
                    ),
                )
                monthly_count += 1
                inserted = True
            explicit_event = clean_text(row.get(mapped["event_type"])) if "event_type" in mapped else None
            status_changed = bool(status and previous_status.get(key) != normalize_header(status))
            if explicit_event or status_changed:
                event_type = explicit_event or status
                conn.execute(
                    """INSERT INTO production_events(source_id,well_key,well_name,event_date,event_type,status,metadata_json)
                       VALUES(?,?,?,?,?,?,?)""",
                    (result.source_id, key, well_name, event_date, event_type, status, Database.json(row)),
                )
                event_count += 1
                inserted = True
            onstream_date = clean_text(row.get(mapped["onstream_date"])) if "onstream_date" in mapped else None
            if onstream_date and key not in onstream_seen:
                conn.execute(
                    """INSERT INTO production_events(source_id,well_key,well_name,event_date,event_type,status,metadata_json)
                       VALUES(?,?,?,?,?,?,?)""",
                    (result.source_id, key, well_name, onstream_date, "投产", status, Database.json(row)),
                )
                onstream_seen.add(key)
                event_count += 1
                inserted = True
            if status:
                previous_status[key] = normalize_header(status)
            if inserted:
                source_rows += 1
            if index % 256 == 0:
                _report_progress(progress_callback, min(0.9, 0.12 + index / 10000), "写入生产动态与井事件")
        result.records = source_rows
        result.details.update({
            "intervals": interval_count, "monthly_records": monthly_count, "events": event_count,
            "field_mapping": mapped, "depth_reference": "MD", "source_system": "OFM-compatible tabular export",
        })

    def _import_attribute_table(self, conn: sqlite3.Connection, path: Path, result: ImportResult, crs: str | None, version: str | None, progress_callback=None) -> None:
        headers, rows = read_tabular(path)
        mapped = field_map(headers)
        if "well" not in mapped:
            raise ValueError("解释/岩心表缺少井名列")
        excluded = {mapped[key] for key in ("well", "uwi", "x", "y", "kb", "td", "top", "base", "version", "batch") if key in mapped}
        attribute_headers = [header for header in headers if header not in excluded]
        seen_wells: set[int] = set()
        attributes_seen: set[str] = set()
        for row_no, row in rows:
            raw_name = clean_text(row.get(mapped["well"]))
            if not raw_name:
                continue
            values = {key: row.get(header) for key, header in mapped.items() if key not in {"well", "top", "base"}}
            values.update({"crs": crs, "attributes": {}})
            well_id = self._resolve_well(conn, raw_name, values.get("uwi"))
            if well_id not in seen_wells:
                self._add_occurrence(conn, result, raw_name, row_no, values)
                seen_wells.add(well_id)
            top = number(row.get(mapped["top"])) if "top" in mapped else None
            base = number(row.get(mapped["base"])) if "base" in mapped else None
            row_version = clean_text(row.get(mapped["version"])) if "version" in mapped else version
            for header in attribute_headers:
                value = clean_text(row.get(header))
                if value is None:
                    continue
                attr_name = str(header).strip()
                attributes_seen.add(attr_name)
                conn.execute(
                    """INSERT INTO interpretation_values(
                       source_id,well_id,row_no,top_md,base_md,attribute_name,attribute_value,numeric_value,version
                       ) VALUES(?,?,?,?,?,?,?,?,?)""",
                    (result.source_id, well_id, row_no, top, base, attr_name, value, number(value), row_version),
                )
                result.records += 1
            if row_no % 256 == 0:
                _report_progress(progress_callback, 0.12 + 0.8 * row_no / max(1, len(rows)), "写入解释 / 岩心属性")
        result.details.update({"wells": len(seen_wells), "attributes": sorted(attributes_seen), "field_mapping": mapped})

    def _import_alias(self, conn: sqlite3.Connection, path: Path, result: ImportResult, **_: Any) -> None:
        headers, rows = read_tabular(path)
        mapped = field_map(headers)
        if not {"alias", "canonical"} <= mapped.keys():
            raise ValueError("别名表需要 ALIAS/别名 和 CANONICAL/标准井名 两列")
        for _, row in rows:
            alias, canonical = clean_text(row.get(mapped["alias"])), clean_text(row.get(mapped["canonical"]))
            if not alias or not canonical:
                continue
            alias_key, canonical_key = normalize_well_name(alias), normalize_well_name(canonical)
            conn.execute(
                "INSERT OR REPLACE INTO alias_rules(alias_key,canonical_key,canonical_name,created_at) VALUES(?,?,?,?)",
                (alias_key, canonical_key, canonical, utcnow()),
            )
            alias_well = conn.execute("SELECT id FROM wells WHERE normalized_key=?", (alias_key,)).fetchone()
            canonical_well = conn.execute("SELECT id FROM wells WHERE normalized_key=?", (canonical_key,)).fetchone()
            if alias_well and canonical_well and alias_well["id"] != canonical_well["id"]:
                self.merge_wells(conn, int(canonical_well["id"]), int(alias_well["id"]), canonical)
            elif alias_well and not canonical_well:
                conn.execute("UPDATE wells SET normalized_key=?,canonical_name=? WHERE id=?", (canonical_key, canonical, alias_well["id"]))
            result.records += 1

    def _import_polygon(self, conn: sqlite3.Connection, path: Path, result: ImportResult, crs: str | None, **_: Any) -> None:
        polygons: list[tuple[str, list]] = []
        if path.suffix.lower() in {".json", ".geojson"}:
            content = json.loads(path.read_text(encoding="utf-8-sig"))
            if content.get("type") == "FeatureCollection":
                features = content.get("features", [])
            elif content.get("type") == "Feature":
                features = [content]
            else:
                features = [{"type": "Feature", "geometry": content, "properties": {}}]
            for index, feature in enumerate(features, 1):
                geometry = feature.get("geometry") or {}
                name = str((feature.get("properties") or {}).get("name") or (feature.get("properties") or {}).get("NAME") or f"Polygon {index}")
                if geometry.get("type") == "Polygon":
                    polygons.append((name, geometry.get("coordinates", [])))
                elif geometry.get("type") == "MultiPolygon":
                    for part_no, coordinates in enumerate(geometry.get("coordinates", []), 1):
                        polygons.append((f"{name} #{part_no}", coordinates))
        else:
            headers, rows = read_tabular(path)
            mapped = field_map(headers)
            if not {"x", "y"} <= mapped.keys():
                raise ValueError("Polygon CSV 至少需要 X、Y 两列")
            grouped: dict[str, list[list[float]]] = defaultdict(list)
            for _, row in rows:
                x, y = number(row.get(mapped["x"])), number(row.get(mapped["y"]))
                if x is None or y is None:
                    continue
                name = clean_text(row.get(mapped["name"])) if "name" in mapped else "Polygon 1"
                grouped[name or "Polygon 1"].append([x, y])
            polygons = [(name, [points]) for name, points in grouped.items()]
        for name, coordinates in polygons:
            if not coordinates or len(coordinates[0]) < 3:
                result.warnings.append(f"{name} 顶点不足 3 个，已跳过")
                continue
            flat = [point for ring in coordinates for point in ring]
            xs, ys = [float(p[0]) for p in flat], [float(p[1]) for p in flat]
            conn.execute(
                """INSERT INTO polygons(source_id,name,crs,geometry_json,area,x_min,x_max,y_min,y_max)
                   VALUES(?,?,?,?,?,?,?,?,?)""",
                (result.source_id, name, crs, Database.json(coordinates), polygon_area(coordinates), min(xs), max(xs), min(ys), max(ys)),
            )
            result.records += 1

    def _import_seismic(self, conn: sqlite3.Connection, path: Path, result: ImportResult, crs: str | None, progress_callback=None, **_: Any) -> None:
        stats = scan_zgy(path) if path.suffix.lower() == ".zgy" else scan_segy(path, lambda value, stage: _report_progress(progress_callback, 0.08 + value * 0.82, stage))
        conn.execute(
            """INSERT INTO seismic_surveys(
               source_id,dimension,trace_count,sample_count_min,sample_count_max,sample_interval_us,format_code,
               x_min,x_max,y_min,y_max,inline_min,inline_max,crossline_min,crossline_max,footprint_area,crs,metadata_json
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                result.source_id, stats["dimension"], stats["trace_count"], stats["sample_count_min"], stats["sample_count_max"],
                stats["sample_interval_us"], stats["format_code"], stats["x_min"], stats["x_max"], stats["y_min"], stats["y_max"],
                stats["inline_min"], stats["inline_max"], stats["crossline_min"], stats["crossline_max"], stats["footprint_area"], crs,
                Database.json({"coordinate_source": stats["coordinate_source"], "text_header_preview": stats["text_header_preview"]}),
            ),
        )
        result.records = stats["trace_count"]
        result.details = stats
        result.warnings.extend(stats.get("warnings", []))

    def merge_wells(self, conn: sqlite3.Connection, keep_id: int, merge_id: int, canonical_name: str | None = None) -> None:
        if keep_id == merge_id:
            return
        keep = conn.execute("SELECT * FROM wells WHERE id=?", (keep_id,)).fetchone()
        merge = conn.execute("SELECT * FROM wells WHERE id=?", (merge_id,)).fetchone()
        if not keep or not merge:
            raise ValueError("待合并井不存在")
        for table in ("well_sources", "las_files", "las_curves", "deviation_stations", "interpretation_values"):
            conn.execute(f"UPDATE {table} SET well_id=? WHERE well_id=?", (keep_id, merge_id))
        conn.execute(
            "INSERT OR REPLACE INTO alias_rules(alias_key,canonical_key,canonical_name,created_at) VALUES(?,?,?,?)",
            (merge["normalized_key"], keep["normalized_key"], canonical_name or keep["canonical_name"], utcnow()),
        )
        conn.execute("DELETE FROM wells WHERE id=?", (merge_id,))
        if canonical_name:
            conn.execute("UPDATE wells SET canonical_name=? WHERE id=?", (canonical_name, keep_id))
        self._refresh_well(conn, keep_id)


def parse_las_header(path: str | Path) -> dict[str, Any]:
    """Read only the LAS header and stop before the ASCII data section.

    This is used for whole-project curve inventory.  It deliberately avoids
    loading multi-megabyte sample arrays; numeric values remain an on-demand
    operation in the distribution tools.
    """
    section = ""
    well: dict[str, dict[str, Any]] = {}
    curves: list[dict[str, Any]] = []
    wrap = False
    with Path(path).open("rb") as handle:
        for raw_line in handle:
            line = None
            # LAS exports from this project contain Spanish Latin-1 accents.
            # Trying GB18030 first can validly but incorrectly turn bytes such
            # as ``é`` into a CJK character, corrupting the mnemonic itself.
            for encoding in ("utf-8-sig", "latin-1", "gb18030"):
                try:
                    line = raw_line.decode(encoding).strip()
                    break
                except UnicodeDecodeError:
                    continue
            if not line or line.startswith("#"):
                continue
            if line.startswith("~"):
                section = line[1:2].upper()
                if section == "A":
                    break
                continue
            if section not in {"V", "W", "C"}:
                continue
            parsed = _parse_las_header_line(line)
            if not parsed:
                continue
            mnemonic = parsed["mnemonic"]
            if section == "V" and mnemonic == "WRAP":
                wrap = str(parsed["value"]).strip().upper() == "YES"
            elif section == "W":
                well[mnemonic] = parsed
            elif section == "C":
                curves.append({
                    **parsed,
                    "sample_count": None,
                    "value_min": None,
                    "value_max": None,
                    "value_mean": None,
                    "value_std": None,
                    "p05": None,
                    "p50": None,
                    "p95": None,
                })
    return {
        "well": well,
        "curves": curves,
        "depth_unit": curves[0].get("unit") if curves else None,
        "start": number(well.get("STRT", {}).get("value")),
        "stop": number(well.get("STOP", {}).get("value")),
        "step": number(well.get("STEP", {}).get("value")),
        "null": number(well.get("NULL", {}).get("value")),
        "wrap": wrap,
        "sample_rows": None,
        "warnings": ["LAS WRAP=YES"] if wrap else [],
    }


def parse_las_curve_samples(
    path: str | Path,
    mnemonics: set[str] | list[str] | tuple[str, ...],
    *,
    sample_limit: int = 4096,
) -> dict[str, Any]:
    """Stream just selected LAS curves for an on-demand chart.

    Unlike :func:`parse_las`, this never loads the whole file nor parses every
    curve column.  It retains a deterministic reservoir of MD/value pairs for
    the requested mnemonics, which is sufficient for an interactive histogram.
    ``WRAP=YES`` exports fall back to the full parser because data row boundaries
    cannot be recovered line by line safely.
    """
    wanted = {str(value or "").upper().strip() for value in mnemonics if str(value or "").strip()}
    if not wanted:
        return {"curves": [], "depth_unit": None, "warnings": []}
    sample_limit = max(128, min(8192, int(sample_limit)))
    section = ""
    well: dict[str, dict[str, Any]] = {}
    all_curves: list[dict[str, Any]] = []
    selected: dict[int, dict[str, Any]] = {}
    expected = 0
    wrap = False
    null_value: float | None = None
    row_count = 0
    md_min: float | None = None
    md_max: float | None = None

    with Path(path).open("rb") as handle:
        for raw_line in handle:
            line = None
            for encoding in ("utf-8-sig", "latin-1", "gb18030"):
                try:
                    line = raw_line.decode(encoding).strip()
                    break
                except UnicodeDecodeError:
                    continue
            if not line or line.startswith("#"):
                continue
            if line.startswith("~"):
                section = line[1:2].upper()
                if section == "A":
                    expected = len(all_curves)
                    null_value = number(well.get("NULL", {}).get("value"))
                    if wrap:
                        break
                    for index, curve in enumerate(all_curves):
                        if curve["mnemonic"] in wanted:
                            selected[index] = {"mnemonic": curve["mnemonic"], "unit": curve.get("unit"), "samples": []}
                continue
            if section in {"V", "W", "C"}:
                parsed = _parse_las_header_line(line)
                if not parsed:
                    continue
                mnemonic = parsed["mnemonic"]
                if section == "V" and mnemonic == "WRAP":
                    wrap = str(parsed["value"]).strip().upper() == "YES"
                elif section == "W":
                    well[mnemonic] = parsed
                elif section == "C":
                    all_curves.append(parsed)
                continue
            if section != "A" or not selected or expected <= 0:
                continue
            values = line.replace(",", " ").split()
            if len(values) < expected:
                continue
            md = number(values[0])
            if md is None:
                continue
            row_count += 1
            md_min = md if md_min is None else min(md_min, md)
            md_max = md if md_max is None else max(md_max, md)
            for index, curve in selected.items():
                value = number(values[index])
                if value is None or (null_value is not None and abs(value - null_value) < 1e-10):
                    continue
                samples = curve["samples"]
                item = {"md": md, "value": value}
                if len(samples) < sample_limit:
                    samples.append(item)
                else:
                    # A stable reservoir avoids retaining all values and avoids
                    # misleading "top-of-well only" distributions.
                    key = f"{Path(path).name}|{curve['mnemonic']}|{row_count}".encode("utf-8", "replace")
                    slot = zlib.crc32(key) % row_count
                    if slot < sample_limit:
                        samples[slot] = item

    if wrap:
        parsed = parse_las(path, include_samples=True, sample_limit=sample_limit)
        curves = [row for row in parsed.get("curves", [])[1:] if str(row.get("mnemonic") or "").upper() in wanted]
        return {"curves": curves, "depth_unit": parsed.get("depth_unit"), "warnings": ["LAS WRAP=YES，使用完整解析"]}
    for curve in selected.values():
        curve["samples"].sort(key=lambda item: item["md"])
    return {
        "curves": list(selected.values()),
        "depth_unit": all_curves[0].get("unit") if all_curves else None,
        "start": number(well.get("STRT", {}).get("value")) or md_min,
        "stop": number(well.get("STOP", {}).get("value")) or md_max,
        "step": number(well.get("STEP", {}).get("value")),
        "sample_rows": row_count,
        "warnings": [],
    }


def parse_las(
    path: str | Path,
    progress_callback: Callable[[float, str], None] | None = None,
    *,
    include_samples: bool = False,
    sample_limit: int = 2048,
) -> dict[str, Any]:
    _report_progress(progress_callback, 0.01, "读取 LAS 文件")
    raw = Path(path).read_bytes()
    text = None
    for encoding in ("utf-8-sig", "latin-1", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("LAS 文件编码无法识别")
    _report_progress(progress_callback, 0.08, "解析 LAS 头段")
    section = ""
    well: dict[str, dict[str, Any]] = {}
    curves: list[dict[str, Any]] = []
    data_lines: list[str] = []
    warnings: list[str] = []
    wrap = False
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("~"):
            section = line[1:2].upper()
            continue
        if section == "A":
            data_lines.append(line)
            continue
        if section not in {"V", "W", "C"}:
            continue
        parsed = _parse_las_header_line(line)
        if not parsed:
            continue
        mnemonic = parsed["mnemonic"]
        if section == "V" and mnemonic == "WRAP":
            wrap = str(parsed["value"]).strip().upper() == "YES"
        elif section == "W":
            well[mnemonic] = parsed
        elif section == "C":
            curves.append({
                **parsed,
                "sample_count": 0,
                "value_min": None,
                "value_max": None,
                "_sum": 0.0,
                "_sum_sq": 0.0,
                "_sample": [],
                "_depth_sample": [],
            })
    null_value = number(well.get("NULL", {}).get("value"))
    expected = len(curves)
    if wrap:
        warnings.append("LAS WRAP=YES：首版按连续数值流重组，建议核对异常行")
        tokens = " ".join(data_lines).replace(",", " ").split()
        data_rows = [tokens[i:i + expected] for i in range(0, len(tokens), expected)] if expected else []
    else:
        data_rows = [line.replace(",", " ").split() for line in data_lines]
    sample_rows = 0
    depth_values: list[float] = []
    total_rows = len(data_rows)
    sample_limit = max(64, min(20_000, int(sample_limit)))
    sample_stride = max(1, math.ceil(total_rows / sample_limit))
    for row_index, values in enumerate(data_rows, 1):
        if len(values) < expected:
            continue
        sample_rows += 1
        depth_value = number(values[0]) if values else None
        keep_depth_sample = include_samples and depth_value is not None and (
            (row_index - 1) % sample_stride == 0 or row_index == total_rows
        )
        for index, curve in enumerate(curves):
            value = number(values[index])
            if value is None or (null_value is not None and abs(value - null_value) < 1e-10):
                continue
            curve["sample_count"] += 1
            curve["_sum"] += value
            curve["_sum_sq"] += value * value
            sample = curve["_sample"]
            if len(sample) < 2048:
                sample.append(value)
            else:
                sample[row_index % 2048] = value
            if keep_depth_sample and index > 0:
                curve["_depth_sample"].append({"md": depth_value, "value": value})
            curve["value_min"] = value if curve["value_min"] is None else min(curve["value_min"], value)
            curve["value_max"] = value if curve["value_max"] is None else max(curve["value_max"], value)
            if index == 0:
                depth_values.append(value)
        if row_index % 4096 == 0:
            _report_progress(progress_callback, 0.12 + 0.84 * row_index / max(1, total_rows), "统计 LAS 样点与分布")
    for curve in curves:
        count = curve["sample_count"]
        if count:
            mean = curve["_sum"] / count
            variance = max(0.0, curve["_sum_sq"] / count - mean * mean)
            ordered = sorted(curve["_sample"])
            curve["value_mean"] = mean
            curve["value_std"] = math.sqrt(variance)
            curve["p05"] = ordered[round((len(ordered) - 1) * 0.05)]
            curve["p50"] = ordered[round((len(ordered) - 1) * 0.50)]
            curve["p95"] = ordered[round((len(ordered) - 1) * 0.95)]
        else:
            curve.update({"value_mean": None, "value_std": None, "p05": None, "p50": None, "p95": None})
        curve.pop("_sum", None)
        curve.pop("_sum_sq", None)
        curve.pop("_sample", None)
        depth_sample = curve.pop("_depth_sample", [])
        if include_samples:
            curve["samples"] = depth_sample
    _report_progress(progress_callback, 1.0, "LAS 解析完成")
    return {
        "well": well,
        "curves": curves,
        "depth_unit": curves[0].get("unit") if curves else None,
        "start": number(well.get("STRT", {}).get("value")) or (min(depth_values) if depth_values else None),
        "stop": number(well.get("STOP", {}).get("value")) or (max(depth_values) if depth_values else None),
        "step": number(well.get("STEP", {}).get("value")),
        "null": null_value,
        "sample_rows": sample_rows,
        "warnings": warnings,
    }


def _parse_las_header_line(line: str) -> dict[str, Any] | None:
    content, _, description = line.partition(":")
    match = re.match(r"^\s*([^\.\s]+)\s*\.([^\s]*)\s*(.*?)\s*$", content)
    if not match:
        return None
    mnemonic, unit, value = match.groups()
    return {"mnemonic": mnemonic.upper(), "unit": unit or None, "value": value.strip(), "description": description.strip() or None}


def scan_zgy(path: str | Path) -> dict[str, Any]:
    """Read ZGY geometry metadata when OpenZGY is available.

    ZGY is indexed even without the optional decoder.  In that case the
    response deliberately contains null geometry instead of guessed cube
    dimensions, and the UI explains how the record was obtained.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    base = {
        "dimension": "3D", "trace_count": 0,
        "sample_count_min": None, "sample_count_max": None,
        "sample_interval_us": None, "format_code": None,
        "x_min": None, "x_max": None, "y_min": None, "y_max": None,
        "inline_min": None, "inline_max": None,
        "crossline_min": None, "crossline_max": None,
        "footprint_area": None, "footprint": [], "coordinate_source": "ZGY 元数据",
        "byte_layout": "zgy", "grid_transform": None, "trace_sample_points": [],
        "text_header_preview": "", "zgy_decoder": False,
        "warnings": ["已识别 ZGY 文件；当前环境没有 OpenZGY 解码组件，内部网格、采样与范围保持待解析。"],
    }
    try:
        from openzgy.api import ZgyReader  # type: ignore
    except (ImportError, ModuleNotFoundError):
        return base
    try:
        with ZgyReader(str(path)) as reader:
            size = tuple(int(value) for value in reader.size)
            annot_start = tuple(float(value) for value in reader.annotstart)
            annot_inc = tuple(float(value) for value in reader.annotinc)
            z_start, z_inc = float(reader.zstart), float(reader.zinc)
            corners = [tuple(float(value) for value in point) for point in reader.corners]
            xs, ys = [point[0] for point in corners], [point[1] for point in corners]
            vertical_dimension = str(getattr(reader, "zunitdim", "") or "").lower()
            base.update({
                "trace_count": size[0] * size[1],
                "sample_count_min": size[2], "sample_count_max": size[2],
                "x_min": min(xs) if xs else None, "x_max": max(xs) if xs else None,
                "y_min": min(ys) if ys else None, "y_max": max(ys) if ys else None,
                "inline_min": annot_start[0], "inline_max": annot_start[0] + (size[0] - 1) * annot_inc[0],
                "crossline_min": annot_start[1], "crossline_max": annot_start[1] + (size[1] - 1) * annot_inc[1],
                "z_start": z_start, "z_increment": z_inc,
                "z_min": min(z_start, z_start + (size[2] - 1) * z_inc),
                "z_max": max(z_start, z_start + (size[2] - 1) * z_inc),
                "z_unit_dimension": vertical_dimension or None,
                "data_type": str(getattr(reader, "datatype", "") or ""),
                "data_range": list(getattr(reader, "datarange", ()) or ()),
                "zgy_decoder": True, "warnings": [],
            })
            return base
    except Exception as exc:
        base["warnings"] = [f"ZGY 元数据读取失败：{exc}"]
        return base


def scan_segy(path: str | Path, progress_callback: Callable[[float, str], None] | None = None) -> dict[str, Any]:
    """Scan SEG-Y textual/binary/trace headers without decoding trace samples.

    The byte-layout detector covers both SEG-Y Rev1 defaults and the common
    Petrel/PEP export where IL/XL occupy 181/185 and scaled CDP X/Y occupy
    189/193. Trace sample blocks are skipped with seek().
    """
    path = Path(path)
    file_size = path.stat().st_size
    if file_size < 3600:
        raise ValueError("文件小于 3600 字节，不是有效 SEG-Y")
    format_sizes = {1: 4, 2: 4, 3: 2, 5: 4, 6: 8, 7: 3, 8: 1, 9: 8, 10: 4, 11: 2, 12: 8, 15: 3, 16: 1}
    # Keep a progressively thinned spatial sample so multi-million-trace files
    # do not require multi-gigabyte memory merely to estimate a footprint.
    spatial_samples: list[tuple[float, float, int | None, int | None]] = []
    max_spatial_samples = 100_000
    sample_stride = 1
    x_min = x_max = y_min = y_max = None
    inline_min = inline_max = crossline_min = crossline_max = None
    sample_count_min = sample_count_max = None
    delay_time_min_ms = delay_time_max_ms = None
    trace_end_time_min_ms = trace_end_time_max_ms = None
    coordinate_source = "none"
    byte_layout = "unknown"
    warnings: list[str] = []
    with path.open("rb") as handle:
        _report_progress(progress_callback, 0.01, "读取 SEG-Y 文本头和二进制头")
        text_header_raw = handle.read(3200)
        text_header = _decode_segy_text_header(text_header_raw)
        binary = handle.read(400)
        sample_interval = struct.unpack(">H", binary[16:18])[0]
        binary_samples = struct.unpack(">H", binary[20:22])[0]
        format_code = struct.unpack(">H", binary[24:26])[0]
        extended_headers = struct.unpack(">h", binary[304:306])[0]
        if extended_headers < 0:
            extended_headers = 0
            warnings.append("扩展文本头数量未知，按 0 处理")
        trace_offset = 3600 + extended_headers * 3200
        if format_code not in format_sizes:
            raise ValueError(f"暂不支持 SEG-Y sample format code {format_code}")
        bytes_per_sample = format_sizes[format_code]
        diagnostic_headers = _sample_trace_headers(
            handle, trace_offset, file_size, binary_samples, bytes_per_sample, 512
        )
        layout = _detect_trace_layout(diagnostic_headers, text_header)
        byte_layout = layout["name"]
        coordinate_source = layout["label"]
        handle.seek(trace_offset)
        trace_count = 0
        while handle.tell() + 240 <= file_size:
            header = handle.read(240)
            if len(header) < 240:
                break
            samples = struct.unpack(">H", header[114:116])[0] or binary_samples
            if samples <= 0:
                warnings.append(f"第 {trace_count + 1} 道样点数为 0，扫描停止")
                break
            x_value, y_value, inline, crossline = _values_from_trace_header(header, layout)
            x, y = x_value, y_value
            if x or y:
                x_min = x_value if x_min is None else min(x_min, x_value)
                x_max = x_value if x_max is None else max(x_max, x_value)
                y_min = y_value if y_min is None else min(y_min, y_value)
                y_max = y_value if y_max is None else max(y_max, y_value)
                if trace_count % sample_stride == 0:
                    spatial_samples.append((x_value, y_value, inline, crossline))
                    if len(spatial_samples) >= max_spatial_samples:
                        spatial_samples = spatial_samples[::2]
                        sample_stride *= 2
            if inline:
                inline_min = inline if inline_min is None else min(inline_min, inline)
                inline_max = inline if inline_max is None else max(inline_max, inline)
            if crossline:
                crossline_min = crossline if crossline_min is None else min(crossline_min, crossline)
                crossline_max = crossline if crossline_max is None else max(crossline_max, crossline)
            sample_count_min = samples if sample_count_min is None else min(sample_count_min, samples)
            sample_count_max = samples if sample_count_max is None else max(sample_count_max, samples)
            delay_time_ms = struct.unpack(">h", header[108:110])[0]
            trace_interval_us = struct.unpack(">H", header[116:118])[0] or sample_interval
            trace_end_time_ms = delay_time_ms + max(0, samples - 1) * trace_interval_us / 1000
            delay_time_min_ms = delay_time_ms if delay_time_min_ms is None else min(delay_time_min_ms, delay_time_ms)
            delay_time_max_ms = delay_time_ms if delay_time_max_ms is None else max(delay_time_max_ms, delay_time_ms)
            trace_end_time_min_ms = trace_end_time_ms if trace_end_time_min_ms is None else min(trace_end_time_min_ms, trace_end_time_ms)
            trace_end_time_max_ms = trace_end_time_ms if trace_end_time_max_ms is None else max(trace_end_time_max_ms, trace_end_time_ms)
            trace_count += 1
            if trace_count % 4096 == 0:
                _report_progress(progress_callback, min(0.98, handle.tell() / file_size), f"扫描 SEG-Y 道头：{trace_count:,} 道")
            next_position = handle.tell() + samples * bytes_per_sample
            if next_position > file_size:
                warnings.append("末道数据长度超出文件边界，文件可能截断")
                break
            handle.seek(next_position)
    if trace_count == 0:
        raise ValueError("未找到有效 SEG-Y trace")
    text_says_2d = bool(re.search(r"(?:ESTUDIO|SURVEY|LINEA|LINE)\s*:?\s*2D|\b2D\b", text_header, re.IGNORECASE))
    text_says_3d = bool(re.search(r"(?:ESTUDIO|SURVEY)\s*:?\s*3D|\b3D\b", text_header, re.IGNORECASE))
    grid_varies = inline_min is not None and inline_max != inline_min and crossline_min is not None and crossline_max != crossline_min
    dimension = "2D" if text_says_2d else ("3D" if text_says_3d or grid_varies else "2D")
    points = [(row[0], row[1]) for row in spatial_samples]
    hull = convex_hull(points)
    if not points:
        warnings.append("未在标准 trace header 位置找到坐标；请确认导出字节位配置")
    elif sample_stride > 1:
        warnings.append(f"道数较多，凸包面积基于每 {sample_stride} 道抽样估算；坐标范围仍为全道扫描")
    grid = _fit_seismic_grid(spatial_samples) if dimension == "3D" else None
    display_trace_samples = []
    if dimension == "2D" and spatial_samples:
        display_stride = max(1, math.ceil(len(spatial_samples) / 5000))
        display_trace_samples = [[round(row[0], 3), round(row[1], 3)] for row in spatial_samples[::display_stride]]
    _report_progress(progress_callback, 1.0, f"SEG-Y 道头扫描完成：{trace_count:,} 道")
    return {
        "dimension": dimension,
        "trace_count": trace_count,
        "sample_count_min": sample_count_min,
        "sample_count_max": sample_count_max,
        "sample_interval_us": sample_interval,
        "delay_time_min_ms": delay_time_min_ms,
        "delay_time_max_ms": delay_time_max_ms,
        "trace_end_time_min_ms": trace_end_time_min_ms,
        "trace_end_time_max_ms": trace_end_time_max_ms,
        "format_code": format_code,
        "x_min": x_min,
        "x_max": x_max,
        "y_min": y_min,
        "y_max": y_max,
        "inline_min": inline_min,
        "inline_max": inline_max,
        "crossline_min": crossline_min,
        "crossline_max": crossline_max,
        "footprint_area": ring_area(hull),
        "footprint": [[round(x, 3), round(y, 3)] for x, y in hull],
        "coordinate_source": coordinate_source,
        "byte_layout": byte_layout,
        "grid_transform": grid,
        "trace_sample_points": display_trace_samples,
        "text_header_preview": text_header[:480].replace("\x00", " ").strip(),
        "warnings": warnings,
    }


def _decode_segy_text_header(raw: bytes) -> str:
    candidates = [raw.decode("ascii", errors="replace"), raw.decode("cp500", errors="replace")]

    def score(text: str) -> float:
        upper = text.upper()
        keywords = sum(upper.count(word) for word in ("C01", "SEG-Y", "SEGY", "INLINE", "XLINE", "SURVEY", "SAMPLE"))
        readable = sum(ch.isalnum() or ch in " .,:;_-/()" for ch in text) / max(1, len(text))
        return keywords * 20 + readable

    return max(candidates, key=score)


def _sample_trace_headers(
    handle: BinaryIO,
    trace_offset: int,
    file_size: int,
    binary_samples: int,
    bytes_per_sample: int,
    limit: int,
) -> list[bytes]:
    headers: list[bytes] = []
    handle.seek(trace_offset)
    while len(headers) < limit and handle.tell() + 240 <= file_size:
        header = handle.read(240)
        if len(header) < 240:
            break
        samples = struct.unpack(">H", header[114:116])[0] or binary_samples
        if samples <= 0:
            break
        headers.append(header)
        next_position = handle.tell() + samples * bytes_per_sample
        if next_position > file_size:
            break
        handle.seek(next_position)
    return headers


def _coordinate_scalar(header: bytes) -> float:
    scalar = struct.unpack(">h", header[70:72])[0]
    return float(scalar) if scalar > 0 else (1.0 / abs(scalar) if scalar < 0 else 1.0)


def _values_from_trace_header(header: bytes, layout: dict[str, Any]) -> tuple[float, float, int | None, int | None]:
    raw_x = struct.unpack(">i", header[layout["x"]:layout["x"] + 4])[0]
    raw_y = struct.unpack(">i", header[layout["y"]:layout["y"] + 4])[0]
    factor = _coordinate_scalar(header) if layout["scaled"] else 1.0
    inline = struct.unpack(">i", header[layout["inline"]:layout["inline"] + 4])[0] if layout.get("inline") is not None else 0
    crossline = struct.unpack(">i", header[layout["crossline"]:layout["crossline"] + 4])[0] if layout.get("crossline") is not None else 0
    return raw_x * factor, raw_y * factor, inline or None, crossline or None


def _detect_trace_layout(headers: list[bytes], text_header: str) -> dict[str, Any]:
    if not headers:
        raise ValueError("SEG-Y 中没有可用于识别字节位的 trace header")
    candidates = [
        {"name": "petrel_pep_swapped", "label": "Scaled CDP X/Y bytes 189/193; IL/XL bytes 181/185", "x": 188, "y": 192, "inline": 180, "crossline": 184, "scaled": True, "prior": 1.0},
        {"name": "segy_rev1", "label": "SEG-Y Rev1 CDP X/Y bytes 181/185; IL/XL bytes 189/193", "x": 180, "y": 184, "inline": 188, "crossline": 192, "scaled": True, "prior": 4.0},
        {"name": "source_xy_petrel_grid", "label": "Unscaled Source X/Y bytes 73/77; IL/XL bytes 181/185", "x": 72, "y": 76, "inline": 180, "crossline": 184, "scaled": False, "prior": 0.0},
        {"name": "source_xy_rev1_grid", "label": "Unscaled Source X/Y bytes 73/77; IL/XL bytes 189/193", "x": 72, "y": 76, "inline": 188, "crossline": 192, "scaled": False, "prior": 0.0},
        {"name": "scaled_source_xy", "label": "Scaled Source X/Y bytes 73/77; IL/XL bytes 189/193", "x": 72, "y": 76, "inline": 188, "crossline": 192, "scaled": True, "prior": 0.0},
        {"name": "source_xy_2d", "label": "Unscaled Source X/Y bytes 73/77; no Inline/Crossline", "x": 72, "y": 76, "inline": None, "crossline": None, "scaled": False, "prior": 5.0 if re.search(r"\b2D\b", text_header, re.IGNORECASE) else -1.0},
    ]
    inline_range = _range_from_text_header(text_header, r"(?:INLINE|IN-LINE|IL)\s*:?\s*(\d+)\s*[-–]\s*(\d+)")
    crossline_range = _range_from_text_header(text_header, r"(?:XLINE|CROSSLINE|CROSS-LINE|XL)\s*:?\s*(\d+)\s*[-–]\s*(\d+)")

    def layout_score(layout: dict[str, Any]) -> float:
        values = [_values_from_trace_header(header, layout) for header in headers]
        coords = [(abs(row[0]), abs(row[1])) for row in values if row[0] or row[1]]
        if not coords:
            return -100.0
        projected = sum(10_000 <= x <= 100_000_000 and 10_000 <= y <= 100_000_000 for x, y in coords) / len(coords)
        geographic = sum(x <= 180 and y <= 90 for x, y in coords) / len(coords)
        score_value = projected * 6 + geographic * 2 + layout.get("prior", 0.0)
        for index, expected in ((2, inline_range), (3, crossline_range)):
            grid_values = [row[index] for row in values if row[index] is not None]
            if not grid_values:
                continue
            if expected:
                margin = max(5, int((expected[1] - expected[0]) * 0.02))
                score_value += 8 * sum(expected[0] - margin <= value <= expected[1] + margin for value in grid_values) / len(grid_values)
            else:
                typical = sum(0 < abs(value) < 10_000_000 for value in grid_values) / len(grid_values)
                score_value += typical
        return score_value

    return max(candidates, key=layout_score)


def _range_from_text_header(text: str, pattern: str) -> tuple[int, int] | None:
    match = re.search(pattern, text, re.IGNORECASE)
    if not match:
        return None
    left, right = int(match.group(1)), int(match.group(2))
    return min(left, right), max(left, right)


def _fit_seismic_grid(samples: list[tuple[float, float, int | None, int | None]]) -> dict[str, Any] | None:
    rows = [(x, y, float(inline), float(crossline)) for x, y, inline, crossline in samples if inline is not None and crossline is not None]
    if len(rows) < 3:
        return None
    # Normal equations for x/y = a*inline + b*crossline + c.
    sii = sum(row[2] * row[2] for row in rows)
    sxx = sum(row[3] * row[3] for row in rows)
    six = sum(row[2] * row[3] for row in rows)
    si = sum(row[2] for row in rows)
    sx = sum(row[3] for row in rows)
    matrix = [[sii, six, si], [six, sxx, sx], [si, sx, float(len(rows))]]

    def fit(target_index: int) -> list[float] | None:
        target = [
            sum(row[2] * row[target_index] for row in rows),
            sum(row[3] * row[target_index] for row in rows),
            sum(row[target_index] for row in rows),
        ]
        return _solve_3x3(matrix, target)

    x_coefficients, y_coefficients = fit(0), fit(1)
    if not x_coefficients or not y_coefficients:
        return None
    squared_error = 0.0
    for x, y, inline, crossline in rows:
        predicted_x = x_coefficients[0] * inline + x_coefficients[1] * crossline + x_coefficients[2]
        predicted_y = y_coefficients[0] * inline + y_coefficients[1] * crossline + y_coefficients[2]
        squared_error += (x - predicted_x) ** 2 + (y - predicted_y) ** 2
    return {
        "x_coefficients": [round(value, 10) for value in x_coefficients],
        "y_coefficients": [round(value, 10) for value in y_coefficients],
        "inline_spacing": round(math.hypot(x_coefficients[0], y_coefficients[0]), 4),
        "crossline_spacing": round(math.hypot(x_coefficients[1], y_coefficients[1]), 4),
        "rms_residual": round(math.sqrt(squared_error / len(rows)), 6),
        "sample_count": len(rows),
    }


def _solve_3x3(matrix: list[list[float]], vector: list[float]) -> list[float] | None:
    augmented = [list(row) + [vector[index]] for index, row in enumerate(matrix)]
    for column in range(3):
        pivot = max(range(column, 3), key=lambda row: abs(augmented[row][column]))
        if abs(augmented[pivot][column]) < 1e-12:
            return None
        augmented[column], augmented[pivot] = augmented[pivot], augmented[column]
        divisor = augmented[column][column]
        augmented[column] = [value / divisor for value in augmented[column]]
        for row in range(3):
            if row == column:
                continue
            factor = augmented[row][column]
            augmented[row] = [augmented[row][index] - factor * augmented[column][index] for index in range(4)]
    return [augmented[row][3] for row in range(3)]


def identity_similarity(a: str, b: str) -> float:
    left, right = normalize_well_name(a), normalize_well_name(b)
    if left == right:
        return 1.0
    ratio = SequenceMatcher(None, left, right).ratio()
    left_numbers = re.findall(r"\d+", left)
    right_numbers = re.findall(r"\d+", right)
    if left_numbers and left_numbers == right_numbers:
        ratio = max(ratio, 0.78)
        if left.startswith(right) or right.startswith(left):
            ratio = max(ratio, 0.9)
    return ratio
